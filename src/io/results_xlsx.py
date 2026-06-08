"""UQ results XLSX workbook export."""

from __future__ import annotations

import json
import math
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

FEATURE_ORDER = [
    "msp", "perplexity", "mte",
    "mcnse", "se-marginal", "se-conditional", "ecc", "eigval", "kle-heat", "kle-matern",
]

REPORT_METRICS = ("auroc", "auprc", "prr", "ece", "brier", "spearman_rho", "kendall_tau")
DEFAULT_REPORT_METRICS = ("auroc",)

FEATURE_GROUP = {
    "msp": "confidence", "perplexity": "confidence", "mte": "confidence",
    "mcnse": "sampling", "se-marginal": "sampling", "se-conditional": "sampling",
    "ecc": "sampling", "eigval": "sampling", "kle-heat": "sampling", "kle-matern": "sampling",
}

DATASET_TITLES = {
    "triviaqa": "TriviaQA",
    "gsm8k": "GSM8K",
    "wmt14_fr_en": "WMT14 fr-en",
    "wmt14_de_en": "WMT14 de-en",
    "xsum": "XSum",
    "samsum": "SamSum",
}

HEADER_FIELDS = [
    ("model_id", "Model"),
    ("gen_length", "Generation length"),
    ("steps", "Steps"),
    ("remasking", "Strategy"),
    ("label_method", "Label method"),
]


@dataclass(frozen=True, order=True)
class CompetitorKey:
    model_id: str
    gen_length: str
    steps: str
    remasking: str
    label_method: str


@dataclass(frozen=True, order=True)
class SummaryColumnKey:
    dataset: str
    model_id: str
    remasking: str


@dataclass
class SummaryColumn:
    key: SummaryColumnKey
    run_count: int = 0
    feature_prr_values: dict[str, list[float]] = field(default_factory=lambda: defaultdict(list))

    def feature_mean_prr(self, feature_name: str) -> float | None:
        vals = self.feature_prr_values.get(feature_name, [])
        return float(sum(vals) / len(vals)) if vals else None


@dataclass
class RunResult:
    run_dir: Path
    metrics_path: Path
    dataset: str
    competitor: CompetitorKey
    qa_accuracy: float | None
    qa_accuracy_weight: int | None
    mean_non_special_tokens: float | None
    token_count_weight: int | None
    feature_metrics: dict[str, dict[str, float]]
    feature_stds: dict[str, dict[str, float]] = field(default_factory=dict)


@dataclass
class CompetitorSummary:
    key: CompetitorKey
    run_dirs: list[Path] = field(default_factory=list)
    qa_accuracy_values: list[tuple[float, int | None]] = field(default_factory=list)
    token_count_values: list[tuple[float, int | None]] = field(default_factory=list)
    feature_values: dict[str, dict[str, list[float]]] = field(
        default_factory=lambda: defaultdict(lambda: defaultdict(list))
    )
    feature_std_values: dict[str, dict[str, list[float]]] = field(
        default_factory=lambda: defaultdict(lambda: defaultdict(list))
    )

    @property
    def run_count(self) -> int:
        return len(self.run_dirs)

    def qa_accuracy(self) -> float | None:
        return _weighted_mean(self.qa_accuracy_values)

    def mean_non_special_tokens(self) -> float | None:
        return _weighted_mean(self.token_count_values)

    def feature_mean(self, feature_name: str, metric_name: str = "auroc") -> float | None:
        vals = self.feature_values.get(feature_name, {}).get(metric_name, [])
        clean = [v for v in vals if math.isfinite(v)]
        return float(sum(clean) / len(clean)) if clean else None

    def feature_std(self, feature_name: str, metric_name: str = "auroc") -> float | None:
        vals = self.feature_std_values.get(feature_name, {}).get(metric_name, [])
        clean = [v for v in vals if math.isfinite(v)]
        return float(sum(clean) / len(clean)) if clean else None


def _numeric(value: Any) -> float | None:
    if value is None:
        return None
    try:
        n = float(value)
    except (TypeError, ValueError):
        return None
    return n if math.isfinite(n) else None


def _weighted_mean(values: Iterable[tuple[float, int | None]]) -> float | None:
    weighted_sum = 0.0
    total_weight = 0
    unweighted = []
    for value, weight in values:
        n = _numeric(value)
        if n is None:
            continue
        unweighted.append(n)
        if weight is not None and int(weight) > 0:
            weighted_sum += n * int(weight)
            total_weight += int(weight)
    if total_weight > 0:
        return float(weighted_sum / total_weight)
    return float(sum(unweighted) / len(unweighted)) if unweighted else None


def _nested_get(payload: dict, keys: tuple[str, ...]) -> Any:
    current: Any = payload
    for key in keys:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def _score_metric_value(score_row: dict, metric_name: str) -> float | None:
    candidates = [
        _nested_get(score_row, ("bootstrap", "metrics", metric_name, "mean")),
        score_row.get(f"{metric_name}_bootstrap_mean"),
        score_row.get(metric_name),
    ]
    for c in candidates:
        n = _numeric(c)
        if n is not None:
            return n
    return None


def _score_metric_std(score_row: dict, metric_name: str) -> float | None:
    candidates = [
        _nested_get(score_row, ("bootstrap", "metrics", metric_name, "std")),
        score_row.get(f"{metric_name}_bootstrap_std"),
    ]
    for c in candidates:
        n = _numeric(c)
        if n is not None:
            return n
    return None


def load_run_result(metrics_path: str | Path) -> RunResult:
    metrics_path = Path(metrics_path).expanduser().resolve()
    run_dir = metrics_path.parent

    with metrics_path.open() as f:
        payload = json.load(f)

    config = {}
    for name in ("config.json", "metadata.json"):
        p = run_dir / name
        if p.exists():
            with p.open() as f:
                data = json.load(f)
            if name == "metadata.json" and isinstance(data.get("run_config"), dict):
                config.update(data["run_config"])
            else:
                config.update(data)

    dataset = str(config.get("dataset", "unknown"))
    model_id = str(config.get("model_id") or config.get("model_name_or_path") or "unknown")
    gen_length = str(config.get("gen_length", "unknown"))
    steps = str(config.get("steps", "unknown"))
    remasking = str(config.get("remasking", "unknown"))
    label_method = str(config.get("label_method", config.get("correctness_method", "unknown")))

    qa_accuracy = _numeric(_nested_get(payload, ("dataset_summary", "accuracy")))
    if qa_accuracy is None:
        qa_accuracy = _numeric(payload.get("accuracy"))
    if qa_accuracy is None:
        nc = _numeric(payload.get("n_correct"))
        ne = _numeric(payload.get("n_examples"))
        if nc is not None and ne is not None and ne > 0:
            qa_accuracy = nc / ne
    qa_weight = _nested_get(payload, ("dataset_summary", "n_labelled"))
    if qa_weight is None:
        qa_weight = _numeric(payload.get("n_examples"))

    mean_non_special = _numeric(payload.get("mean_visible_tokens"))

    feature_metrics: dict[str, dict[str, float]] = {}
    feature_stds: dict[str, dict[str, float]] = {}
    for score in payload.get("scores", []) or []:
        if not isinstance(score, dict):
            continue
        name = score.get("score_name")
        if not name:
            continue
        values = {}
        stds = {}
        for m in REPORT_METRICS:
            v = _score_metric_value(score, m)
            if v is not None:
                values[m] = v
            s = _score_metric_std(score, m)
            if s is not None:
                stds[m] = s
        if values:
            feature_metrics[str(name)] = values
        if stds:
            feature_stds[str(name)] = stds

    for name, feat_data in (payload.get("features") or {}).items():
        if not isinstance(feat_data, dict) or name in feature_metrics:
            continue
        values = {}
        stds = {}
        for m in REPORT_METRICS:
            v = _score_metric_value(feat_data, m)
            if v is not None:
                values[m] = v
            s = _score_metric_std(feat_data, m)
            if s is not None:
                stds[m] = s
        if values:
            feature_metrics[str(name)] = values
        if stds:
            feature_stds[str(name)] = stds

    return RunResult(
        run_dir=run_dir,
        metrics_path=metrics_path,
        dataset=dataset,
        competitor=CompetitorKey(
            model_id=model_id, gen_length=gen_length, steps=steps,
            remasking=remasking, label_method=label_method,
        ),
        qa_accuracy=qa_accuracy,
        qa_accuracy_weight=int(qa_weight) if qa_weight is not None else None,
        mean_non_special_tokens=mean_non_special,
        token_count_weight=int(qa_weight) if qa_weight is not None else None,
        feature_metrics=feature_metrics,
        feature_stds=feature_stds,
    )


def find_metrics_files(
    roots: Iterable[str | Path],
    metrics_name: str = "uq_eval_metrics.json",
) -> list[Path]:
    files: list[Path] = []
    seen: set[Path] = set()
    skip_names = {"__pycache__", ".git", ".pytest_cache"}
    for root in roots:
        root = Path(root).expanduser()
        if not root.exists():
            continue
        candidates = [root] if root.is_file() else root.rglob(metrics_name)
        for c in candidates:
            if c.is_file() and c.name == metrics_name:
                if not any(p in skip_names for p in c.parts):
                    resolved = c.resolve()
                    if resolved not in seen:
                        files.append(resolved)
                        seen.add(resolved)
    return sorted(files)


def find_kfold_summaries(
    roots: Iterable[str | Path],
    summary_name: str = "kfold_summary.json",
) -> list[Path]:
    files: list[Path] = []
    seen: set[Path] = set()
    for root in roots:
        root = Path(root).expanduser()
        if not root.exists():
            continue
        for c in root.rglob(summary_name):
            if c.is_file():
                resolved = c.resolve()
                if resolved not in seen:
                    files.append(resolved)
                    seen.add(resolved)
    return sorted(files)


def _load_config_from_ancestors(start_dir: Path) -> dict:
    config: dict = {}
    for name in ("config.json", "metadata.json"):
        for d in [start_dir, *start_dir.parents]:
            p = d / name
            if p.exists():
                with p.open() as f:
                    data = json.load(f)
                if name == "metadata.json" and isinstance(data.get("run_config"), dict):
                    config.update(data["run_config"])
                else:
                    config.update(data)
                break
    return config


def _aggregate_kfold_fold_stats(kfold_dir: Path, n_folds: int) -> tuple[float | None, int | None, float | None]:
    total_correct = 0
    total_examples = 0
    mvt_sum = 0.0
    mvt_count = 0
    for i in range(n_folds):
        p = kfold_dir / f"fold_{i}" / "baseline" / "uq_eval_metrics.json"
        if not p.exists():
            continue
        with p.open() as f:
            fold_data = json.load(f)
        nc = _numeric(fold_data.get("n_correct"))
        ne = _numeric(fold_data.get("n_examples"))
        if nc is not None and ne is not None:
            total_correct += int(nc)
            total_examples += int(ne)
        mvt = _numeric(fold_data.get("mean_visible_tokens"))
        ne_int = int(ne) if ne is not None else 0
        if mvt is not None and ne_int > 0:
            mvt_sum += mvt * ne_int
            mvt_count += ne_int

    qa_accuracy = total_correct / total_examples if total_examples > 0 else None
    mean_visible = mvt_sum / mvt_count if mvt_count > 0 else None
    weight = total_examples if total_examples > 0 else None
    return qa_accuracy, weight, mean_visible


def load_kfold_results(summary_path: str | Path) -> list[RunResult]:
    summary_path = Path(summary_path).expanduser().resolve()
    kfold_dir = summary_path.parent
    run_dir = kfold_dir.parent

    with summary_path.open() as f:
        payload = json.load(f)

    config = _load_config_from_ancestors(run_dir)
    dataset = str(config.get("dataset", "unknown"))
    model_id = str(config.get("model_id") or config.get("model_name_or_path") or "unknown")
    gen_length = str(config.get("gen_length", "unknown"))
    steps = str(config.get("steps", "unknown"))
    remasking = str(config.get("remasking", "unknown"))
    label_method = str(config.get("label_method", config.get("correctness_method", "unknown")))

    competitor = CompetitorKey(
        model_id=model_id, gen_length=gen_length, steps=steps,
        remasking=remasking, label_method=label_method,
    )

    n_folds = payload.get("n_folds", 0)
    qa_accuracy, qa_weight, mean_visible = _aggregate_kfold_fold_stats(kfold_dir, n_folds)

    results: list[RunResult] = []
    for section, subdir_name in [("baseline", "baseline"), ("selection", "selection")]:
        section_data = payload.get(section)
        if not section_data or not isinstance(section_data, dict):
            continue
        feature_metrics: dict[str, dict[str, float]] = {}
        feature_stds: dict[str, dict[str, float]] = {}
        for feat, metric_data in section_data.items():
            if not isinstance(metric_data, dict):
                continue
            values: dict[str, float] = {}
            stds: dict[str, float] = {}
            for m in REPORT_METRICS:
                entry = metric_data.get(m)
                if isinstance(entry, dict):
                    n = _numeric(entry.get("mean"))
                    s = _numeric(entry.get("std"))
                else:
                    n = _numeric(entry)
                    s = None
                if n is not None:
                    values[m] = n
                if s is not None:
                    stds[m] = s
            if values:
                feature_metrics[str(feat)] = values
            if stds:
                feature_stds[str(feat)] = stds
        if feature_metrics:
            results.append(RunResult(
                run_dir=kfold_dir / subdir_name,
                metrics_path=summary_path,
                dataset=dataset,
                competitor=competitor,
                qa_accuracy=qa_accuracy,
                qa_accuracy_weight=qa_weight,
                mean_non_special_tokens=mean_visible,
                token_count_weight=qa_weight,
                feature_metrics=feature_metrics,
                feature_stds=feature_stds,
            ))
    return results


def aggregate_results(results: Iterable[RunResult]) -> dict[str, list[CompetitorSummary]]:
    by_ds: dict[str, dict[CompetitorKey, CompetitorSummary]] = defaultdict(dict)
    for r in results:
        title = DATASET_TITLES.get(r.dataset, r.dataset)
        bucket = by_ds[title]
        s = bucket.setdefault(r.competitor, CompetitorSummary(key=r.competitor))
        s.run_dirs.append(r.run_dir)
        if r.qa_accuracy is not None:
            s.qa_accuracy_values.append((r.qa_accuracy, r.qa_accuracy_weight))
        is_selection = r.run_dir.name == "selection"
        for feat, metrics in r.feature_metrics.items():
            key_name = _selection_feature_name(feat) if is_selection else feat
            for m, v in metrics.items():
                s.feature_values[key_name][m].append(float(v))
        for feat, stds in r.feature_stds.items():
            key_name = _selection_feature_name(feat) if is_selection else feat
            for m, v in stds.items():
                s.feature_std_values[key_name][m].append(float(v))
    return {ds: sorted(b.values(), key=lambda x: x.key) for ds, b in sorted(by_ds.items())}


def aggregate_summary_columns(dataset_summaries: dict[str, list[CompetitorSummary]]) -> list[SummaryColumn]:
    grouped: dict[SummaryColumnKey, SummaryColumn] = {}
    for dataset, summaries in dataset_summaries.items():
        for s in summaries:
            key = SummaryColumnKey(dataset=dataset, model_id=s.key.model_id, remasking=s.key.remasking)
            col = grouped.setdefault(key, SummaryColumn(key=key))
            col.run_count += s.run_count
            for feat, metrics in s.feature_values.items():
                for v in metrics.get("prr", []):
                    n = _numeric(v)
                    if n is not None:
                        col.feature_prr_values[feat].append(n)
    return sorted(grouped.values(), key=lambda c: (c.key.dataset, c.key.model_id, c.key.remasking))


def _feature_order(features: Iterable[str]) -> list[str]:
    seen = set(features)
    ordered = [f for f in FEATURE_ORDER if f in seen]
    ordered.extend(sorted(f for f in seen if f not in FEATURE_ORDER))
    return ordered


def _feature_group(name: str) -> str:
    if name in FEATURE_GROUP:
        return FEATURE_GROUP[name]
    if name.startswith(("selected-", "full-", "random-")):
        return "trajectory"
    if name.startswith("kle-"):
        return "sampling"
    return "other"


def _performance_fill(value: float, min_val: float, max_val: float, PatternFill):
    ratio = (value - min_val) / (max_val - min_val) if max_val > min_val else 1.0
    ratio = max(0.0, min(1.0, ratio))
    red = (244, 204, 204)
    green = (217, 234, 211)
    rgb = tuple(round(red[i] + (green[i] - red[i]) * ratio) for i in range(3))
    return PatternFill("solid", fgColor=f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")


def _normalize_report_metrics(metrics: Iterable[str] | None) -> list[str]:
    values: list[str] = []
    for item in metrics or DEFAULT_REPORT_METRICS:
        for part in str(item).split(","):
            m = part.strip().lower()
            if m:
                values.append(m)
    values = list(dict.fromkeys(values)) or list(DEFAULT_REPORT_METRICS)
    invalid = [m for m in values if m not in REPORT_METRICS]
    if invalid:
        raise ValueError(f"Unsupported report metric(s): {', '.join(invalid)}")
    return values


def _fmt_mean_std(mean: float | None, std: float | None) -> str | None:
    if mean is None:
        return None
    if std is not None:
        return f"{mean:.3f} ({std:.3f})"
    return f"{mean:.3f}"


def write_workbook(
    dataset_summaries: dict[str, list[CompetitorSummary]],
    output_path: str | Path,
    *,
    report_metrics: Iterable[str] | None = None,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    report_metrics_list = _normalize_report_metrics(report_metrics)
    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    subheader_fill = PatternFill("solid", fgColor="D9EAF7")
    qa_fill = PatternFill("solid", fgColor="E2F0D9")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    if not dataset_summaries:
        ws = wb.create_sheet("Summary")
        ws["A1"] = "No evaluated runs found"
        ws["A1"].font = Font(bold=True)
    else:
        for dataset, summaries in dataset_summaries.items():
            title = re.sub(r"[\\/*?:\[\]]", "_", dataset)[:31]
            ws = wb.create_sheet(title)

            metric_columns = [(s, m) for s in summaries for m in report_metrics_list]
            ws.column_dimensions["A"].width = 28
            ws.sheet_view.showGridLines = False

            for row_idx, (_, label) in enumerate(HEADER_FIELDS, start=1):
                cell = ws.cell(row=row_idx, column=1, value=label)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = left
                cell.border = border

            metric_row = len(HEADER_FIELDS) + 1
            ws.cell(row=metric_row, column=1, value="Metric").font = Font(bold=True, color="FFFFFF")
            ws.cell(row=metric_row, column=1).fill = header_fill
            ws.cell(row=metric_row, column=1).alignment = left
            ws.cell(row=metric_row, column=1).border = border

            for col_off, (s, m) in enumerate(metric_columns, start=2):
                vals = {
                    "model_id": s.key.model_id, "gen_length": s.key.gen_length,
                    "steps": s.key.steps, "remasking": s.key.remasking,
                    "label_method": s.key.label_method,
                }
                for row_idx, (field_name, _) in enumerate(HEADER_FIELDS, start=1):
                    cell = ws.cell(row=row_idx, column=col_off, value=vals[field_name])
                    cell.font = Font(bold=True, color="FFFFFF" if row_idx == 1 else "000000")
                    cell.fill = header_fill if row_idx == 1 else subheader_fill
                    cell.alignment = center
                    cell.border = border

                cell = ws.cell(row=metric_row, column=col_off, value=m.upper())
                cell.font = Font(bold=True)
                cell.fill = subheader_fill
                cell.alignment = center
                cell.border = border
                ws.column_dimensions[get_column_letter(col_off)].width = 16

            meta_start = metric_row + 1
            for row_idx, (label, getter, fill, fmt) in enumerate([
                ("Runs included", None, subheader_fill, "0"),
                ("Greedy accuracy", "qa_accuracy", qa_fill, "0.0%"),
            ], start=meta_start):
                ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
                ws.cell(row=row_idx, column=1).fill = fill
                ws.cell(row=row_idx, column=1).alignment = left
                ws.cell(row=row_idx, column=1).border = border
                for col_off, (s, _) in enumerate(metric_columns, start=2):
                    v = s.run_count if getter is None else getattr(s, getter)()
                    cell = ws.cell(row=row_idx, column=col_off, value=v)
                    cell.fill = fill
                    cell.alignment = center
                    cell.border = border
                    cell.number_format = fmt

            feat_header_row = meta_start + 3
            ws.freeze_panes = f"B{feat_header_row + 1}"
            ws.cell(row=feat_header_row, column=1, value="UQ feature").font = Font(bold=True, color="FFFFFF")
            ws.cell(row=feat_header_row, column=1).fill = header_fill
            ws.cell(row=feat_header_row, column=1).alignment = left
            ws.cell(row=feat_header_row, column=1).border = border

            feature_names = _feature_order(
                name for s in summaries for name in s.feature_values
            )

            all_vals = []
            for s in summaries:
                for fn in feature_names:
                    for m in report_metrics_list:
                        v = s.feature_mean(fn, m)
                        if v is not None:
                            all_vals.append(v)
            min_v = min(all_vals) if all_vals else 0.0
            max_v = max(all_vals) if all_vals else 0.0

            prev_group = None
            medium = Side(style="medium", color="7F7F7F")
            group_border = Border(left=thin, right=thin, top=medium, bottom=thin)
            for row_off, fn in enumerate(feature_names, start=feat_header_row + 1):
                group = _feature_group(fn)
                row_b = group_border if prev_group is not None and group != prev_group else border
                prev_group = group
                ws.cell(row=row_off, column=1, value=fn).alignment = left
                ws.cell(row=row_off, column=1).border = row_b
                for col_off, (s, m) in enumerate(metric_columns, start=2):
                    v = s.feature_mean(fn, m)
                    std = s.feature_std(fn, m)
                    label = _fmt_mean_std(v, std) if v is not None else "N/A"
                    cell = ws.cell(row=row_off, column=col_off, value=label)
                    cell.alignment = center
                    cell.border = row_b
                    if v is not None:
                        cell.fill = _performance_fill(v, min_v, max_v, PatternFill)

    output_path = Path(output_path).expanduser()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


# ---------------------------------------------------------------------------
# New model-centric XLSX layout
# ---------------------------------------------------------------------------

from src.config import MODEL_HF_IDS, Model, Dataset, Remasking

HF_ID_TO_MODEL: dict[str, str] = {v: k.value for k, v in MODEL_HF_IDS.items()}

MODEL_ORDER = [m.value for m in Model]
DATASET_ORDER = [d.value for d in Dataset]
REMASKING_ORDER = [r.value for r in Remasking]

DATASET_SHORT = {
    "triviaqa": "TriviaQA",
    "gsm8k": "GSM8K",
    "wmt14_fr_en": "WMT fr-en",
    "wmt14_de_en": "WMT de-en",
    "xsum": "XSum",
    "samsum": "SamSum",
}

_SAMPLING_FEATURES = ["se-marginal", "se-conditional", "ecc", "eigval", "kle-heat", "kle-matern"]
FEATURE_GROUPS: list[tuple[str, list[str]]] = [
    ("Token UQ", ["msp", "perplexity", "mte"]),
    ("IID Sample UQ", ["mcnse", "se-marginal", "se-conditional", "ecc", "eigval", "kle-heat", "kle-matern"]),
    ("Full Trajectory UQ", [f"full-{f}" for f in _SAMPLING_FEATURES]),
    ("Selected Step UQ", [f"selected-{f}" for f in _SAMPLING_FEATURES]),
    ("Random Scope UQ", [f"random-{f}" for f in _SAMPLING_FEATURES]),
    ("Trajectory AD", ["full-ad", "weighted-ad"]),
]


def _selection_feature_name(feat: str) -> str:
    if feat == "ad":
        return "weighted-ad"
    return f"selected-{feat}"


def _resolve_model_name(model_id: str) -> str:
    if model_id in HF_ID_TO_MODEL:
        return HF_ID_TO_MODEL[model_id]
    for hf_id, name in HF_ID_TO_MODEL.items():
        if name.lower() in model_id.lower() or model_id.endswith(hf_id.split("/")[-1]):
            return name
    return model_id


def _sort_key_dataset_remasking(key: tuple[str, str]) -> tuple[int, int]:
    ds, rm = key
    ds_idx = DATASET_ORDER.index(ds) if ds in DATASET_ORDER else len(DATASET_ORDER)
    rm_idx = REMASKING_ORDER.index(rm) if rm in REMASKING_ORDER else len(REMASKING_ORDER)
    return (ds_idx, rm_idx)


def aggregate_by_model(
    results: Iterable[RunResult],
) -> dict[str, dict[tuple[str, str], CompetitorSummary]]:
    by_model: dict[str, dict[tuple[str, str], CompetitorSummary]] = {}
    for r in results:
        model_name = _resolve_model_name(r.competitor.model_id)
        col_key = (r.dataset, r.competitor.remasking)
        bucket = by_model.setdefault(model_name, {})
        s = bucket.setdefault(col_key, CompetitorSummary(key=r.competitor))
        s.run_dirs.append(r.run_dir)
        if r.qa_accuracy is not None:
            s.qa_accuracy_values.append((r.qa_accuracy, r.qa_accuracy_weight))
        if r.mean_non_special_tokens is not None:
            s.token_count_values.append((r.mean_non_special_tokens, r.token_count_weight))
        is_selection = r.run_dir.name == "selection"
        for feat, metrics in r.feature_metrics.items():
            key_name = _selection_feature_name(feat) if is_selection else feat
            for m, v in metrics.items():
                s.feature_values[key_name][m].append(float(v))
        for feat, stds in r.feature_stds.items():
            key_name = _selection_feature_name(feat) if is_selection else feat
            for m, v in stds.items():
                s.feature_std_values[key_name][m].append(float(v))
    return by_model


def write_model_workbook(
    results: Iterable[RunResult],
    output_path: str | Path,
    *,
    metric: str = "prr",
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    by_model = aggregate_by_model(results)
    if not by_model:
        return

    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    subheader_fill = PatternFill("solid", fgColor="D9EAF7")
    group_fill = PatternFill("solid", fgColor="E8E8E8")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    bold_white = Font(bold=True, color="FFFFFF", size=9)
    bold_black = Font(bold=True, size=9)
    normal = Font(size=9)

    for model_name in MODEL_ORDER:
        if model_name not in by_model:
            continue
        columns = by_model[model_name]
        col_keys = sorted(columns.keys(), key=_sort_key_dataset_remasking)
        if not col_keys:
            continue

        ws = wb.create_sheet(model_name)
        ws.sheet_view.showGridLines = False

        # -- Row 1: dataset names (merged across rd/lc pairs) --
        ws.cell(row=1, column=1).fill = header_fill
        ws.cell(row=1, column=1).border = border
        ws.column_dimensions["A"].width = 22

        col = 2
        prev_ds = None
        merge_start = None
        for ds, rm in col_keys:
            ws.column_dimensions[get_column_letter(col)].width = 16
            if ds != prev_ds:
                if merge_start is not None and col > merge_start + 1:
                    ws.merge_cells(start_row=1, start_column=merge_start, end_row=1, end_column=col - 1)
                merge_start = col
                prev_ds = ds
            col += 1
        if merge_start is not None and col > merge_start + 1:
            ws.merge_cells(start_row=1, start_column=merge_start, end_row=1, end_column=col - 1)

        col = 2
        prev_ds = None
        for ds, rm in col_keys:
            if ds != prev_ds:
                cell = ws.cell(row=1, column=col, value=DATASET_SHORT.get(ds, ds))
                cell.font = bold_white
                cell.fill = header_fill
                cell.alignment = center
                cell.border = border
                prev_ds = ds
            else:
                ws.cell(row=1, column=col).fill = header_fill
                ws.cell(row=1, column=col).border = border
            col += 1

        # -- Row 2: remasking strategy --
        ws.cell(row=2, column=1, value="Strategy").font = bold_white
        ws.cell(row=2, column=1).fill = header_fill
        ws.cell(row=2, column=1).alignment = left
        ws.cell(row=2, column=1).border = border
        for i, (ds, rm) in enumerate(col_keys, start=2):
            cell = ws.cell(row=2, column=i, value=rm)
            cell.font = bold_black
            cell.fill = subheader_fill
            cell.alignment = center
            cell.border = border

        # -- Row 3: Accuracy --
        row = 3
        ws.cell(row=row, column=1, value="Accuracy").font = bold_black
        ws.cell(row=row, column=1).alignment = left
        ws.cell(row=row, column=1).border = border
        for i, key in enumerate(col_keys, start=2):
            s = columns[key]
            v = s.qa_accuracy()
            cell = ws.cell(row=row, column=i, value=v)
            cell.number_format = "0.0%" if v is not None else "@"
            cell.font = normal
            cell.alignment = center
            cell.border = border

        # -- Row 4: Mean non-special tokens --
        row = 4
        ws.cell(row=row, column=1, value="Mean tokens").font = bold_black
        ws.cell(row=row, column=1).alignment = left
        ws.cell(row=row, column=1).border = border
        for i, key in enumerate(col_keys, start=2):
            s = columns[key]
            v = s.mean_non_special_tokens()
            cell = ws.cell(row=row, column=i, value=round(v, 1) if v is not None else None)
            cell.number_format = "0.0" if v is not None else "@"
            cell.font = normal
            cell.alignment = center
            cell.border = border

        # -- Collect all metric values for heatmap scaling --
        all_vals = []
        for group_label, features in FEATURE_GROUPS:
            for feat in features:
                for key in col_keys:
                    s = columns[key]
                    v = s.feature_mean(feat, metric)
                    if v is not None:
                        all_vals.append(v)
        min_v = min(all_vals) if all_vals else 0.0
        max_v = max(all_vals) if all_vals else 0.0

        # -- Feature rows grouped by category --
        row = 6
        for group_label, features in FEATURE_GROUPS:
            present = [f for f in features if any(columns[k].feature_mean(f, metric) is not None for k in col_keys)]
            if not present:
                continue

            ws.cell(row=row, column=1, value=group_label).font = bold_black
            ws.cell(row=row, column=1).fill = group_fill
            ws.cell(row=row, column=1).alignment = left
            ws.cell(row=row, column=1).border = border
            for i in range(2, len(col_keys) + 2):
                ws.cell(row=row, column=i).fill = group_fill
                ws.cell(row=row, column=i).border = border
            row += 1

            for feat in present:
                ws.cell(row=row, column=1, value=feat).font = normal
                ws.cell(row=row, column=1).alignment = left
                ws.cell(row=row, column=1).border = border
                for i, key in enumerate(col_keys, start=2):
                    s = columns[key]
                    v = s.feature_mean(feat, metric)
                    std = s.feature_std(feat, metric)
                    label = _fmt_mean_std(v, std)
                    cell = ws.cell(row=row, column=i, value=label)
                    cell.font = normal
                    cell.alignment = center
                    cell.border = border
                    if v is not None and max_v > min_v:
                        cell.fill = _performance_fill(v, min_v, max_v, PatternFill)
                row += 1

        # -- Metric label at bottom --
        row += 1
        ws.cell(row=row, column=1, value=f"Metric: {metric.upper()}").font = Font(italic=True, size=8, color="666666")

        ws.freeze_panes = "B3"

    output_path = Path(output_path).expanduser()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"[report] Wrote XLSX to {output_path}")
