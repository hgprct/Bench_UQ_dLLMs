"""Tests for the XLSX export module."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from src.utils.results_xlsx import (
    CompetitorKey,
    CompetitorSummary,
    RunResult,
    aggregate_results,
    aggregate_summary_columns,
    find_metrics_files,
    load_run_result,
)


class TestCompetitorKey:
    def test_ordering(self):
        k1 = CompetitorKey("Dream", "128", "64", "lc", "exact_match")
        k2 = CompetitorKey("LLaDA", "128", "64", "lc", "exact_match")
        assert k1 < k2


class TestAggregateResults:
    def test_groups_by_dataset(self):
        r1 = RunResult(
            run_dir=Path("/a"),
            metrics_path=Path("/a/m.json"),
            dataset="triviaqa",
            competitor=CompetitorKey("LLaDA", "128", "64", "lc", "exact_match"),
            qa_accuracy=0.8,
            qa_accuracy_weight=100,
            mean_non_special_tokens=None,
            token_count_weight=None,
            feature_metrics={"msp": {"auroc": 0.7}},
        )
        r2 = RunResult(
            run_dir=Path("/b"),
            metrics_path=Path("/b/m.json"),
            dataset="gsm8k",
            competitor=CompetitorKey("Dream", "64", "32", "rd", "exact_match"),
            qa_accuracy=0.5,
            qa_accuracy_weight=50,
            mean_non_special_tokens=None,
            token_count_weight=None,
            feature_metrics={"msp": {"auroc": 0.6}},
        )
        result = aggregate_results([r1, r2])
        assert "TriviaQA" in result
        assert "GSM8K" in result
        assert len(result["TriviaQA"]) == 1
        assert result["TriviaQA"][0].qa_accuracy() == pytest.approx(0.8)


class TestFindMetricsFiles:
    def test_finds_json_files(self, tmp_path):
        run_dir = tmp_path / "outputs" / "run1"
        run_dir.mkdir(parents=True)
        metrics = run_dir / "uq_eval_metrics.json"
        metrics.write_text("{}")
        files = find_metrics_files([tmp_path / "outputs"])
        assert len(files) == 1
        assert files[0].name == "uq_eval_metrics.json"

    def test_skips_missing_roots(self):
        files = find_metrics_files([Path("/nonexistent/path")])
        assert files == []


class TestSummaryColumns:
    def test_aggregates_prr(self):
        summary = CompetitorSummary(
            key=CompetitorKey("LLaDA", "128", "64", "lc", "exact_match"),
        )
        summary.feature_values["msp"]["prr"].extend([0.3, 0.5])
        result = aggregate_summary_columns({"TriviaQA": [summary]})
        assert len(result) == 1
        assert result[0].feature_mean_prr("msp") == pytest.approx(0.4)


class TestAggregateByModel:
    def test_groups_by_model_and_config(self):
        from src.utils.results_xlsx import aggregate_by_model

        r1 = RunResult(
            run_dir=Path("/a"),
            metrics_path=Path("/a/m.json"),
            dataset="triviaqa",
            competitor=CompetitorKey("GSAI-ML/LLaDA-8B-Instruct", "128", "64", "lc", "llm_judge"),
            qa_accuracy=0.8,
            qa_accuracy_weight=100,
            mean_non_special_tokens=45.0,
            token_count_weight=100,
            feature_metrics={"msp": {"prr": 0.32}},
        )
        r2 = RunResult(
            run_dir=Path("/b"),
            metrics_path=Path("/b/m.json"),
            dataset="triviaqa",
            competitor=CompetitorKey("GSAI-ML/LLaDA-8B-Instruct", "128", "64", "rd", "llm_judge"),
            qa_accuracy=0.7,
            qa_accuracy_weight=100,
            mean_non_special_tokens=42.0,
            token_count_weight=100,
            feature_metrics={"msp": {"prr": 0.28}},
        )
        r3 = RunResult(
            run_dir=Path("/c"),
            metrics_path=Path("/c/m.json"),
            dataset="gsm8k",
            competitor=CompetitorKey("Dream-org/Dream-v0-Instruct-7B", "64", "32", "lc", "exact_match"),
            qa_accuracy=0.5,
            qa_accuracy_weight=50,
            mean_non_special_tokens=30.0,
            token_count_weight=50,
            feature_metrics={"msp": {"prr": 0.45}},
        )
        by_model = aggregate_by_model([r1, r2, r3])
        assert "LLaDA" in by_model
        assert "Dream" in by_model
        assert ("triviaqa", "lc") in by_model["LLaDA"]
        assert ("triviaqa", "rd") in by_model["LLaDA"]
        assert ("gsm8k", "lc") in by_model["Dream"]
        assert by_model["LLaDA"][("triviaqa", "lc")].qa_accuracy() == pytest.approx(0.8)
        assert by_model["LLaDA"][("triviaqa", "lc")].feature_mean("msp", "prr") == pytest.approx(0.32)


class TestWriteModelWorkbook:
    def test_creates_xlsx_with_model_tabs(self, tmp_path):
        from src.utils.results_xlsx import write_model_workbook

        results = [
            RunResult(
                run_dir=Path("/a"),
                metrics_path=Path("/a/m.json"),
                dataset="triviaqa",
                competitor=CompetitorKey("GSAI-ML/LLaDA-8B-Instruct", "128", "64", "lc", "llm_judge"),
                qa_accuracy=0.8,
                qa_accuracy_weight=100,
                mean_non_special_tokens=45.0,
                token_count_weight=100,
                feature_metrics={"msp": {"prr": 0.32}, "perplexity": {"prr": 0.28}},
            ),
            RunResult(
                run_dir=Path("/b"),
                metrics_path=Path("/b/m.json"),
                dataset="triviaqa",
                competitor=CompetitorKey("Dream-org/Dream-v0-Instruct-7B", "64", "32", "rd", "exact_match"),
                qa_accuracy=0.5,
                qa_accuracy_weight=50,
                mean_non_special_tokens=30.0,
                token_count_weight=50,
                feature_metrics={"msp": {"prr": 0.45}},
            ),
        ]
        out_path = tmp_path / "test.xlsx"
        write_model_workbook(results, out_path, metric="prr")
        assert out_path.exists()

        from openpyxl import load_workbook
        wb = load_workbook(out_path)
        assert "LLaDA" in wb.sheetnames
        assert "Dream" in wb.sheetnames
